from flask import Flask, jsonify, request, render_template, send_file
from flask_cors import CORS
from database import create_tables, get_db_connection

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO

app = Flask(__name__)
CORS(app)

create_tables()


@app.route("/")
def home():
    return render_template("index.html")

@app.route("/clients")
def clients_page():
    return render_template("clients.html")

@app.route("/meetings")
def meetings_page():
    return render_template("meetings.html")


# Get all clients
@app.route("/api/clients", methods=["GET"])
def get_clients():
    connection = get_db_connection()

    clients = connection.execute(
        "SELECT * FROM clients ORDER BY id DESC"
    ).fetchall()

    connection.close()

    return jsonify([dict(client) for client in clients])


# Add a new client
@app.route("/api/clients", methods=["POST"])
def add_client():
    data = request.get_json()

    name = data.get("name")
    company = data.get("company")
    country = data.get("country")
    email = data.get("email")
    phone = data.get("phone")
    status = data.get("status", "Active")

    if not name or not company:
        return jsonify({
            "error": "Name and company are required"
        }), 400

    connection = get_db_connection()

    cursor = connection.execute("""
        INSERT INTO clients
        (name, company, country, email, phone, status)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (name, company, country, email, phone, status))

    connection.commit()

    client_id = cursor.lastrowid

    connection.close()

    return jsonify({
        "message": "Client added successfully",
        "id": client_id
    }), 201


# Delete a client
@app.route("/api/clients/<int:client_id>", methods=["DELETE"])
def delete_client(client_id):
    connection = get_db_connection()

    connection.execute(
        "DELETE FROM clients WHERE id = ?",
        (client_id,)
    )

    connection.commit()
    connection.close()

    return jsonify({
        "message": "Client deleted successfully"
    })

# Get all meetings
@app.route("/api/meetings", methods=["GET"])
def get_meetings():
    connection = get_db_connection()

    meetings = connection.execute("""
        SELECT meetings.*, clients.name AS client_name,
               clients.company AS client_company
        FROM meetings
        LEFT JOIN clients ON meetings.client_id = clients.id
        ORDER BY meetings.date ASC, meetings.time ASC
    """).fetchall()

    connection.close()

    return jsonify([dict(meeting) for meeting in meetings])


# Add a new meeting
@app.route("/api/meetings", methods=["POST"])
def add_meeting():
    data = request.get_json()

    client_id = data.get("client_id")
    title = data.get("title")
    date = data.get("date")
    time = data.get("time")
    meeting_type = data.get("meeting_type")
    participants = data.get("participants")
    agenda = data.get("agenda")
    status = data.get("status", "Upcoming")

    if not client_id or not title or not date or not time:
        return jsonify({
            "error": "Client, title, date and time are required"
        }), 400

    connection = get_db_connection()

    cursor = connection.execute("""
        INSERT INTO meetings
        (client_id, title, date, time, meeting_type,
         participants, agenda, status)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        client_id,
        title,
        date,
        time,
        meeting_type,
        participants,
        agenda,
        status
    ))

    connection.commit()

    meeting_id = cursor.lastrowid

    connection.close()

    return jsonify({
        "message": "Meeting scheduled successfully",
        "id": meeting_id
    }), 201


# Delete a meeting
@app.route("/api/meetings/<int:meeting_id>", methods=["DELETE"])
def delete_meeting(meeting_id):
    connection = get_db_connection()

    connection.execute(
        "DELETE FROM meetings WHERE id = ?",
        (meeting_id,)
    )

    connection.commit()
    connection.close()

    return jsonify({
        "message": "Meeting deleted successfully"
    })

@app.route("/tasks")
def tasks_page():
    return render_template("tasks.html")

@app.route("/reports")
def reports_page():
    return render_template("reports.html")


@app.route("/api/tasks", methods=["GET"])
def get_tasks():
    connection = get_db_connection()

    tasks = connection.execute("""
        SELECT * FROM tasks
        ORDER BY due_date ASC
    """).fetchall()

    connection.close()

    return jsonify([dict(task) for task in tasks])


@app.route("/api/tasks", methods=["POST"])
def add_task():
    data = request.get_json()

    title = data.get("title")
    description = data.get("description")
    assigned_to = data.get("assigned_to")
    due_date = data.get("due_date")
    priority = data.get("priority", "Medium")
    status = data.get("status", "Pending")

    if not title:
        return jsonify({
            "error": "Task title is required"
        }), 400

    connection = get_db_connection()

    cursor = connection.execute("""
        INSERT INTO tasks
        (title, description, assigned_to, due_date, priority, status)
        VALUES (?, ?, ?, ?, ?, ?)
    """, (
        title,
        description,
        assigned_to,
        due_date,
        priority,
        status
    ))

    connection.commit()

    task_id = cursor.lastrowid

    connection.close()

    return jsonify({
        "message": "Task added successfully",
        "id": task_id
    }), 201


@app.route("/api/tasks/<int:task_id>", methods=["PUT"])
def update_task(task_id):
    data = request.get_json()

    status = data.get("status")

    if not status:
        return jsonify({
            "error": "Status is required"
        }), 400

    connection = get_db_connection()

    connection.execute("""
        UPDATE tasks
        SET status = ?
        WHERE id = ?
    """, (status, task_id))

    connection.commit()
    connection.close()

    return jsonify({
        "message": "Task updated successfully"
    })


@app.route("/api/tasks/<int:task_id>", methods=["DELETE"])
def delete_task(task_id):
    connection = get_db_connection()

    connection.execute(
        "DELETE FROM tasks WHERE id = ?",
        (task_id,)
    )

    connection.commit()
    connection.close()

    return jsonify({
        "message": "Task deleted successfully"
    })

# Export clients to Excel
@app.route("/export/clients")
def export_clients():
    connection = get_db_connection()

    clients = connection.execute("""
        SELECT id, name, company, country, email, phone, status
        FROM clients
        ORDER BY id DESC
    """).fetchall()

    connection.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clients"

    headers = [
        "ID",
        "Name",
        "Company",
        "Country",
        "Email",
        "Phone",
        "Status"
    ]

    sheet.append(headers)

    for client in clients:
        sheet.append([
            client["id"],
            client["name"],
            client["company"],
            client["country"],
            client["email"],
            client["phone"],
            client["status"]
        ])

    # Format header
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Add filter
    sheet.auto_filter.ref = sheet.dimensions

    # Freeze header row
    sheet.freeze_panes = "A2"

    # Adjust column widths
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        sheet.column_dimensions[column_letter].width = max_length + 3

    file = BytesIO()
    workbook.save(file)
    file.seek(0)

    return send_file(
        file,
        as_attachment=True,
        download_name="OfficeFlow_Clients.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# Export meetings to Excel
@app.route("/export/meetings")
def export_meetings():
    connection = get_db_connection()

    meetings = connection.execute("""
        SELECT meetings.*, 
               clients.name AS client_name,
               clients.company AS client_company
        FROM meetings
        LEFT JOIN clients ON meetings.client_id = clients.id
        ORDER BY meetings.date ASC, meetings.time ASC
    """).fetchall()

    connection.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Meetings"

    headers = [
        "ID",
        "Client",
        "Company",
        "Title",
        "Date",
        "Time",
        "Meeting Type",
        "Participants",
        "Agenda",
        "Status"
    ]

    sheet.append(headers)

    for meeting in meetings:
        sheet.append([
            meeting["id"],
            meeting["client_name"],
            meeting["client_company"],
            meeting["title"],
            meeting["date"],
            meeting["time"],
            meeting["meeting_type"],
            meeting["participants"],
            meeting["agenda"],
            meeting["status"]
        ])

    # Format header
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Add filter
    sheet.auto_filter.ref = sheet.dimensions

    # Freeze header row
    sheet.freeze_panes = "A2"

    # Adjust column widths
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        sheet.column_dimensions[column_letter].width = max_length + 3

    file = BytesIO()
    workbook.save(file)
    file.seek(0)

    return send_file(
        file,
        as_attachment=True,
        download_name="OfficeFlow_Meetings.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# Export tasks to Excel
@app.route("/export/tasks")
def export_tasks():
    connection = get_db_connection()

    tasks = connection.execute("""
        SELECT id, title, description, assigned_to,
               due_date, priority, status
        FROM tasks
        ORDER BY due_date ASC
    """).fetchall()

    connection.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tasks"

    headers = [
        "ID",
        "Title",
        "Description",
        "Assigned To",
        "Due Date",
        "Priority",
        "Status"
    ]

    sheet.append(headers)

    for task in tasks:
        sheet.append([
            task["id"],
            task["title"],
            task["description"],
            task["assigned_to"],
            task["due_date"],
            task["priority"],
            task["status"]
        ])

    # Format header
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Add filter
    sheet.auto_filter.ref = sheet.dimensions

    # Freeze header row
    sheet.freeze_panes = "A2"

    # Adjust column widths
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        sheet.column_dimensions[column_letter].width = max_length + 3

    file = BytesIO()
    workbook.save(file)
    file.seek(0)

    return send_file(
        file,
        as_attachment=True,
        download_name="OfficeFlow_Tasks.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route("/export/administrative-report")
def export_administrative_report():

    # Create workbook
    workbook = Workbook()

    # Connect to database
    conn = get_db_connection()

    clients = conn.execute("SELECT * FROM clients").fetchall()
    meetings = conn.execute("SELECT * FROM meetings").fetchall()
    tasks = conn.execute("SELECT * FROM tasks").fetchall()

    conn.close()


    # =========================
    # SUMMARY SHEET
    # =========================

    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    summary_sheet.append([
        "OfficeFlow Administrative Report",
        ""
    ])

    summary_sheet.append([
        "",
        ""
    ])

    summary_sheet.append([
        "Category",
        "Total"
    ])

    summary_sheet.append([
        "Total Clients",
        len(clients)
    ])

    summary_sheet.append([
        "Total Meetings",
        len(meetings)
    ])

    summary_sheet.append([
        "Total Tasks",
        len(tasks)
    ])


    # Format summary
    summary_sheet["A1"].font = Font(
        bold=True,
        size=16
    )

    for cell in summary_sheet[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center"
        )

    summary_sheet.column_dimensions["A"].width = 30
    summary_sheet.column_dimensions["B"].width = 20


    # =========================
    # FUNCTION TO CREATE SHEET
    # =========================

    def create_data_sheet(workbook, sheet_name, rows):

        sheet = workbook.create_sheet(sheet_name)

        if not rows:
            sheet.append(["No data available"])
            return

        # Get actual column names from database
        headers = rows[0].keys()

        # Add headers
        sheet.append(list(headers))

        # Add database rows
        for row in rows:
            sheet.append([
                row[column]
                for column in headers
            ])

        # Format header
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center"
            )

        # Freeze header
        sheet.freeze_panes = "A2"

        # Add filter
        sheet.auto_filter.ref = sheet.dimensions

        # Automatic column width
        for column in sheet.columns:

            max_length = 0

            column_letter = column[0].column_letter

            for cell in column:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 3,
                40
            )


    # =========================
    # CREATE DATA SHEETS
    # =========================

    create_data_sheet(
        workbook,
        "Clients",
        clients
    )

    create_data_sheet(
        workbook,
        "Meetings",
        meetings
    )

    create_data_sheet(
        workbook,
        "Tasks",
        tasks
    )


    # =========================
    # CREATE EXCEL FILE
    # =========================

    file = BytesIO()

    workbook.save(file)

    file.seek(0)

    return send_file(
        file,
        as_attachment=True,
        download_name="OfficeFlow_Administrative_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

if __name__ == "__main__":
    app.run(debug=True)